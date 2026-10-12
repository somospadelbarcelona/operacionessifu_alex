# -*- coding: utf-8 -*-
"""
SIFU ULTRA-FAST PYTHON EXTRACTOR v5.1 (GIRONA)
Extrae instantaneamente MASTER GENERAL GIRONA OKK.xlsx hacia master_data.js.
Compatible con archivos abiertos en Microsoft Excel usando CreateFileW con modo compartido total.
"""

import os
import sys
import time
import io
import json
from datetime import datetime

try:
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8', line_buffering=True, errors='replace')
    if hasattr(sys.stderr, 'reconfigure'):
        sys.stderr.reconfigure(encoding='utf-8', line_buffering=True, errors='replace')
except Exception:
    pass

# Directorio base
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, "MASTER GENERAL GIRONA OKK.xlsx")
OUTPUT_JS = os.path.join(BASE_DIR, "master_data.js")
OUTPUT_JSON = os.path.join(BASE_DIR, "girona_data.json")

def read_bytes_shared(filepath, max_retries=10, delay=0.3):
    """
    Lee los bytes del archivo en Windows incluso si está abierto en Excel,
    usando FILE_SHARE_READ | FILE_SHARE_WRITE | FILE_SHARE_DELETE (0x7).
    Esto evita por completo el error [Errno 13] Permission denied.
    """
    if not os.path.exists(filepath):
        return None

    for attempt in range(max_retries):
        try:
            if os.name == 'nt':
                import ctypes
                import msvcrt

                kernel32 = ctypes.windll.kernel32
                GENERIC_READ = 0x80000000
                FILE_SHARE_READ = 1
                FILE_SHARE_WRITE = 2
                FILE_SHARE_DELETE = 4
                OPEN_EXISTING = 3
                FILE_ATTRIBUTE_NORMAL = 0x80

                handle = kernel32.CreateFileW(
                    filepath,
                    GENERIC_READ,
                    FILE_SHARE_READ | FILE_SHARE_WRITE | FILE_SHARE_DELETE,
                    None,
                    OPEN_EXISTING,
                    FILE_ATTRIBUTE_NORMAL,
                    None
                )

                if handle != -1 and handle != 0:
                    fd = msvcrt.open_osfhandle(handle, os.O_RDONLY)
                    with open(fd, 'rb', closefd=True) as f:
                        data = f.read()
                    if len(data) > 0:
                        return data
            else:
                with open(filepath, 'rb') as f:
                    data = f.read()
                if len(data) > 0:
                    return data
        except Exception:
            pass

        time.sleep(delay)

    return None

def extract(excel_file=EXCEL_PATH):
    if not os.path.exists(excel_file):
        print(f"[ERROR] No se encuentra el archivo: {excel_file}", flush=True)
        return False

    import openpyxl

    file_bytes = read_bytes_shared(excel_file, max_retries=10, delay=0.3)
    if not file_bytes:
        print(f"[ERROR] No se pudo leer el archivo tras 10 intentos (Excel bloqueado durante guardado)", flush=True)
        return False

    wb = None
    try:
        buf = io.BytesIO(file_bytes)
        wb = openpyxl.load_workbook(buf, read_only=True, data_only=True)

        # 1. Extraer Hoja Principal (SEGUIMIENTO)
        sheet_name = 'SEGUIMIENTO' if 'SEGUIMIENTO' in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]

        headers = []
        rows = []
        max_valid_col = 24  # Columnas útiles hasta ob1

        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                seen = {}
                for idx, h in enumerate(row):
                    h_str = str(h or '').strip()
                    if h_str in ('ESTADO.1', 'ESTADO.11') or h_str.startswith('ESTADO.'):
                        h_str = 'ESTADO1'
                    if not h_str:
                        h_str = f'Columna{idx+1}'
                    h_upper = h_str.upper()
                    if h_upper in seen:
                        count = seen[h_upper]
                        seen[h_upper] += 1
                        h_str = f'{h_str}{count}'
                    else:
                        seen[h_upper] = 1
                    headers.append(h_str)

                for idx in range(len(headers)-1, -1, -1):
                    if not headers[idx].startswith('Columna'):
                        max_valid_col = max(max_valid_col, idx + 1)
                        break
            else:
                if not any(row):
                    continue

                row_vals = [str(c or '').strip() for c in row[:max_valid_col]]
                non_empty = [v for v in row_vals if v]
                if len(non_empty) <= 1 and (len(non_empty) == 0 or len(non_empty[0]) <= 1):
                    continue

                row_dict = {}
                for col_idx in range(max_valid_col):
                    h_name = headers[col_idx] if col_idx < len(headers) else f'Columna{col_idx+1}'
                    val = row[col_idx] if col_idx < len(row) else ''
                    if val is None:
                        val_str = ''
                    elif isinstance(val, datetime):
                        val_str = val.strftime('%d/%m/%Y')
                    else:
                        val_str = str(val).strip()
                    row_dict[h_name] = val_str

                rows.append(row_dict)

        # 2. Extraer Hoja DESCUBIERTOS (si existe)
        descubiertos_rows = []
        if 'DESCUBIERTOS' in wb.sheetnames:
            ws_desc = wb['DESCUBIERTOS']
            desc_headers = []
            for i, row in enumerate(ws_desc.iter_rows(values_only=True)):
                if i == 0:
                    for idx, h in enumerate(row[:max_valid_col]):
                        desc_headers.append(str(h or f'Columna{idx+1}').strip())
                else:
                    non_empty = [c for c in row[:max_valid_col] if c is not None and str(c).strip()]
                    if len(non_empty) <= 1:
                        continue
                    row_dict = {}
                    for col_idx in range(len(desc_headers)):
                        val = row[col_idx] if col_idx < len(row) else ''
                        row_dict[desc_headers[col_idx]] = str(val or '').strip()
                    descubiertos_rows.append(row_dict)

        wb.close()

        # Generar contenido JS
        now_str = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        json_data = json.dumps(rows, ensure_ascii=False)
        json_desc = json.dumps(descubiertos_rows, ensure_ascii=False)

        js_content = (
            f"// Ultima actualizacion automatica: {now_str}\n"
            f"window.MASTER_DATA_TIMESTAMP = '{now_str}';\n"
            f"window.INITIAL_MASTER_DATA = {json_data};\n"
            f"window.INITIAL_DESCUBIERTOS_DATA = {json_desc};\n"
        )

        with open(OUTPUT_JS, 'w', encoding='utf-8') as f:
            f.write(js_content)

        with open(OUTPUT_JSON, 'w', encoding='utf-8') as f:
            json.dump({"timestamp": now_str, "services": rows, "descubiertos": descubiertos_rows}, f, ensure_ascii=False, indent=2)

        print(f"[OK] {now_str} - Sincronizado: {len(rows)} servicios ({len(descubiertos_rows)} descubiertos) en master_data.js", flush=True)
        return True

    except Exception as e:
        print(f"[ERROR] Error al procesar Excel: {e}", flush=True)
        try:
            if wb: wb.close()
        except Exception:
            pass
        return False

if __name__ == "__main__":
    t0 = time.time()
    ok = extract()
    t1 = time.time()
    if ok:
        print(f"[TIEMPO] Extraccion completada en {t1-t0:.2f} segundos.", flush=True)
    sys.exit(0 if ok else 1)
