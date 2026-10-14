# -*- coding: utf-8 -*-
"""
Constructor del Censo Maestro de Trabajadores de SIFU Girona.
Combina y cruza:
1. ho_employees_cache.json (empleados oficiales de SAP S171 Girona)
2. MASTER GENERAL GIRONA OKK.xlsx (titulares y suplentes con servicios)
3. Censo previo de fichajes_sync_gcalidad_app.py
4. Overrides manuales en fichajes_censo_manual.json
Genera fichajes_censo_girona.json deduplicado y normalizado.
"""

import os
import json
import re
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
HO_EMPLOYEES_FILE = os.path.join(BASE_DIR, "ho_employees_cache.json")
EXCEL_FILE = os.path.join(BASE_DIR, "MASTER GENERAL GIRONA OKK.xlsx")
MANUAL_FILE = os.path.join(BASE_DIR, "fichajes_censo_manual.json")
OUTPUT_FILE = os.path.join(BASE_DIR, "fichajes_censo_girona.json")

def normalize_dni(dni):
    if not dni:
        return ""
    dni = re.sub(r'[\s\-_.]', '', str(dni)).upper()
    return dni

def normalize_name(name):
    if not name:
        return ""
    name = re.sub(r'\s+', ' ', str(name)).strip().upper()
    return name

def build_census():
    census_by_dni = {}
    name_to_dni = {}

    # 1. Cargar censo base de fichajes_sync_gcalidad_app.py si existe
    prior_census = []
    py_sync = os.path.join(BASE_DIR, "fichajes_sync_gcalidad_app.py")
    if os.path.exists(py_sync):
        with open(py_sync, "r", encoding="utf-8") as f:
            content = f.read()
        m = re.search(r'GIRONA_CENSUS\s*=\s*(\[.*?\]\n)', content, re.DOTALL)
        if m:
            import ast
            try:
                prior_census = ast.literal_eval(m.group(1))
            except Exception as e:
                print("Error parsing prior census:", e)

    for item in prior_census:
        dni = normalize_dni(item.get("dni"))
        name = normalize_name(item.get("name"))
        if dni:
            census_by_dni[dni] = {
                "dni": dni,
                "name": name,
                "servicio": item.get("servicio", "GIRONA GENERAL"),
                "sector": item.get("sector", "AUXILIARES"),
                "id_empleado": "",
                "telefono": "",
                "email": "",
                "fuente": "censo_previo"
            }
            if name:
                name_to_dni[name] = dni

    # 2. Cargar ho_employees_cache.json (Centro S171 Girona)
    if os.path.exists(HO_EMPLOYEES_FILE):
        try:
            with open(HO_EMPLOYEES_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            by_name = data.get("by_name", {})
            by_dni = data.get("by_dni", {})

            for d, info in by_dni.items():
                dni = normalize_dni(d)
                if not dni or not isinstance(info, dict):
                    continue
                name = normalize_name(info.get("nombre_completo"))
                if dni in census_by_dni:
                    # Actualizar con datos enriquecidos de SAP
                    census_by_dni[dni]["id_empleado"] = info.get("id_empleado", "")
                    census_by_dni[dni]["telefono"] = info.get("telefono", "")
                    census_by_dni[dni]["email"] = info.get("email", "")
                    if not census_by_dni[dni]["name"]:
                        census_by_dni[dni]["name"] = name
                else:
                    census_by_dni[dni] = {
                        "dni": dni,
                        "name": name,
                        "servicio": "CENTRO GIRONA S171",
                        "sector": "OPERACIONES",
                        "id_empleado": info.get("id_empleado", ""),
                        "telefono": info.get("telefono", ""),
                        "email": info.get("email", ""),
                        "fuente": "sap_ho_cache"
                    }
                if name:
                    name_to_dni[name] = dni
        except Exception as e:
            print("Error cargando ho_employees_cache.json:", e)

    # 3. Cruzar con MASTER GENERAL GIRONA OKK.xlsx para asociar servicio/sector exactos
    if os.path.exists(EXCEL_FILE):
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
            for sheetname in ["SEGUIMIENTO", "DESCUBIERTOS"]:
                if sheetname in wb.sheetnames:
                    ws = wb[sheetname]
                    # Asumimos fila 1 cabecera
                    headers = [str(cell.value).strip().upper() if cell.value is not None else "" for cell in ws[1]]
                    
                    def find_col(possible_names):
                        for idx, h in enumerate(headers):
                            for pn in possible_names:
                                if pn in h:
                                    return idx
                        return -1

                    col_tipo = find_col(["TIPO S", "TIPO SERVICIO"])
                    col_serv = find_col(["SERVICIO", "N1 OK"])
                    col_titular = find_col(["TITULAR"])
                    col_suplente = find_col(["SUPLENTE", "CUBRE"])

                    for row in ws.iter_rows(min_row=2, values_only=True):
                        tipo_s = str(row[col_tipo]).strip() if col_tipo >= 0 and row[col_tipo] else ""
                        serv = str(row[col_serv]).strip() if col_serv >= 0 and row[col_serv] else ""
                        titular = normalize_name(row[col_titular]) if col_titular >= 0 and row[col_titular] else ""
                        suplente = normalize_name(row[col_suplente]) if col_suplente >= 0 and row[col_suplente] else ""

                        for person in [titular, suplente]:
                            if not person:
                                continue
                            # Buscar si tenemos DNI de esta persona
                            found_dni = name_to_dni.get(person)
                            if not found_dni:
                                # Búsqueda parcial difusa
                                for n_k, d_v in name_to_dni.items():
                                    if person in n_k or n_k in person:
                                        found_dni = d_v
                                        break
                            if found_dni and found_dni in census_by_dni:
                                if serv and (census_by_dni[found_dni]["servicio"] in ["GIRONA GENERAL", "CENTRO GIRONA S171"] or not census_by_dni[found_dni]["servicio"]):
                                    census_by_dni[found_dni]["servicio"] = serv
                                if tipo_s and (census_by_dni[found_dni]["sector"] in ["AUXILIARES", "OPERACIONES"] or not census_by_dni[found_dni]["sector"]):
                                    census_by_dni[found_dni]["sector"] = tipo_s
        except Exception as e:
            print("Error cruzando con Excel maestro:", e)

    # 4. Cargar manual overrides si existen
    if os.path.exists(MANUAL_FILE):
        try:
            with open(MANUAL_FILE, "r", encoding="utf-8") as f:
                manual = json.load(f)
            for m_item in manual:
                dni = normalize_dni(m_item.get("dni"))
                if not dni:
                    continue
                if dni in census_by_dni:
                    census_by_dni[dni].update(m_item)
                else:
                    census_by_dni[dni] = m_item
                    census_by_dni[dni]["fuente"] = "manual"
        except Exception as e:
            print("Error cargando manual overrides:", e)

    # Excluir explícitamente al personal de Jardinería (solicitud expresa de Operaciones)
    DNIS_JARDINERIA = {
        'X1368710A', '40437050Z', '40345869M', '40437178G', '40441302B',
        '41559902Y', '23274238D', '41551869T', '40433702R', '40443871G',
        '77606815P', '40441589C', '42389334Z', '41538356Z', '40336214C',
        '46798028Z', '40325492J', '40348705C', '30588235H', '40321262F'
    }
    APELLIDOS_JARDINERIA = [
        'BEN MOUSSA', 'SANCHEZ PEREZ', 'OSORIO VACA', 'CIFUENTES BOSCH',
        'SALVADOR VILANOVA', 'ESPINAS DEL FRESNO', 'FELIU LOPEZ',
        'MARTI BARTIS', 'BELLAPART', 'ALDAMIZ', 'SIMO CASTELLS'
    ]

    census_list = []
    for w in census_by_dni.values():
        sec_serv = (str(w.get('sector', '')) + ' ' + str(w.get('servicio', ''))).upper()
        if 'JARDIN' in sec_serv:
            continue
        if w.get('dni', '').upper() in DNIS_JARDINERIA:
            continue
        w_name = w.get('name', '').upper()
        if any(ap in w_name for ap in APELLIDOS_JARDINERIA):
            continue
        census_list.append(w)

    # Ordenar por nombre
    census_list = sorted(census_list, key=lambda x: x.get("name", ""))

    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(census_list, f, ensure_ascii=False, indent=2)

    print(f"Censo Maestro generado con exito (excluida Jardineria): {len(census_list)} trabajadores en {OUTPUT_FILE}")
    return census_list

if __name__ == "__main__":
    build_census()
