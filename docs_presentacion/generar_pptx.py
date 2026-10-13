from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
GRIS = RGBColor(0x33, 0x33, 0x33)

slides = [
    ("El problema: la información está repartida",
     ["Datos en el Excel maestro, fichajes y programas de personal (HO y CMH)",
      "Hay que abrir varios sitios para saber qué pasa hoy",
      "Los descubiertos y las ausencias se detectan tarde"],
     "Pregunta al público: ¿cuántos Excel o programas abrís cada mañana?"),
    ("La solución: un único panel alimentado por nuestro Excel",
     ["Seguimos editando el Excel maestro como siempre",
      "El panel lee los datos y los muestra al instante",
      "También cruza fichajes y datos de personal"],
     "Insiste en que no cambia nuestra forma de trabajar. Solo se añade una pantalla que lo muestra todo junto."),
    ("Cómo funciona: guardas el Excel y la pantalla se actualiza",
     ["1. Editas y guardas el Excel (Ctrl+S)",
      "2. Un proceso en segundo plano detecta el cambio",
      "3. El panel se refresca solo, sin recargar la página"],
     "Aquí va la demo en vivo: guarda un cambio y enséñalo en pantalla."),
    ("Qué podemos ver: el día a día en un sitio",
     ["Panel general: estado de los servicios del día",
      "Fichajes: quién ha fichado y quién no",
      "Supervisor virtual: revisa incidencias automáticamente",
      "Pedidos, vacaciones, bajas y sustituciones",
      "Calidad, contratos e informes exportables",
      "Avisos y chat interno"],
     "Enseña 2 o 3 pantallas, no todas. Elige las que más usa tu equipo."),
    ("Beneficios y siguientes pasos",
     ["Detección rápida de descubiertos y ausencias",
      "Una única fuente de información",
      "Funciona en el móvil y se puede instalar como app",
      "Siguiente paso: recoger vuestras sugerencias de mejora"],
     "Cierra con: ¿qué os gustaría ver en el panel?"),
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Portada
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Panel de Control Operativo SIFU Girona"
s.placeholders[1].text = "Todo el día a día, en una sola pantalla"
for p in s.shapes.title.text_frame.paragraphs:
    for r in p.runs:
        r.font.color.rgb = AZUL
        r.font.bold = True

for titulo, puntos, notas in slides:
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = titulo
    for p in s.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(36)
            r.font.bold = True
            r.font.color.rgb = AZUL
    tf = s.placeholders[1].text_frame
    tf.text = puntos[0]
    for t in puntos[1:]:
        tf.add_paragraph().text = t
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(26)
            r.font.color.rgb = GRIS
    s.notes_slide.notes_text_frame.text = notas

prs.save("docs_presentacion/Presentacion_Panel_SIFU_Girona.pptx")
print("OK")
