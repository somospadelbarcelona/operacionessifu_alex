# -*- coding: utf-8 -*-
import sys
try:
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

import os
import json
import re
import unicodedata
from datetime import datetime
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, "MASTER GENERAL GIRONA OKK.xlsx")

from ho_service import HOService, RESPONSABLE_ID

def normalize_text(text):
    if not text:
        return ""
    n = unicodedata.normalize('NFKD', str(text)).encode('ASCII', 'ignore').decode('utf-8')
    n = re.sub(r'[^A-Z0-9\s]', ' ', n.upper())
    return " ".join(n.split())

def parse_sap_date(d_str):
    if not d_str:
        return "Indefinido"
    try:
        ts = int(d_str.replace('/Date(', '').replace(')/', '')) / 1000
        return datetime.fromtimestamp(ts).strftime('%d/%m/%Y')
    except Exception:
        return str(d_str).strip()

ho = HOService()
ho.connect()

print("1. Descargando proyectos de SAP HO...")
ho_projects_list = ho.get_girona_projects(max_results=500)
ho_projects = {p.get('IdProyecto'): p for p in ho_projects_list if p.get('IdProyecto')}

print("2. Descargando empleados de Girona de SAP HO...")
ho_employees = {}
for c in ['S171', 'S172', 'B171']:
    res = ho.query_entity_set("EmpleadoSet", filter_query=f"IdCentro eq '{c}'", top=500)
    for e in res or []:
        emp_id = e.get('IdEmpleado')
        nom_norm = normalize_text(e.get('NombreCompleto'))
        e_processed = {
            "id": emp_id,
            "nombre": (e.get('NombreCompleto') or '').strip(),
            "dni": (e.get('DNI') or '').strip(),
            "telefono": (e.get('Telefono') or '').strip(),
            "centro": e.get('IdCentro'),
            "tipo_contrato": e.get('DenTipoContrato') or 'Fijo',
            "fin_contrato": parse_sap_date(e.get('FinContrato')),
            "vac_restantes": int(e.get('vacRestantes') or 0),
            "vac_tomados": int(e.get('vacTomados') or 0),
            "vac_totales": int(e.get('vacTotales') or 0),
            "horas_semana": e.get('HorasSemana'),
            "status": e.get('NomStatus')
        }
        ho_employees[emp_id] = e_processed
        if nom_norm:
            ho_employees[nom_norm] = e_processed

print("3. Descargando asignaciones oficiales en SAP HO...")
ho_asig_by_project = {}
for proy_id in ho_projects.keys():
    asigs = ho.query_entity_set("EmpleadoAsigProySet", filter_query=f"IdProyecto eq '{proy_id}'")
    if asigs:
        ho_asig_by_project[proy_id] = [
            {"id": a.get('IdEmpleado'), "nombre": (a.get('Nombre') or '').strip(), "dni": (a.get('DNI') or '').strip()}
            for a in asigs
        ]

print("4. Analizando MASTER GENERAL GIRONA OKK.xlsx...")
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb['SEGUIMIENTO']

excel_rows = []
for r in range(2, ws.max_row + 1):
    proy = str(ws.cell(r, 6).value or '').strip()
    tit = str(ws.cell(r, 8).value or '').strip()
    serv = str(ws.cell(r, 4).value or '').strip()
    tipo_s = str(ws.cell(r, 2).value or '').strip()
    tipo_serv = str(ws.cell(r, 3).value or '').strip()
    dir_val = str(ws.cell(r, 5).value or '').strip()
    estado = str(ws.cell(r, 7).value or '').strip()
    fin_con = str(ws.cell(r, 13).value or '').strip()
    num_vac = str(ws.cell(r, 14).value or '').strip()
    
    if not any([proy, tit, serv, tipo_s]):
        continue

    excel_rows.append({
        "row": r,
        "proyecto": proy,
        "titular": tit,
        "servicio": serv,
        "tipo_s": tipo_s,
        "tipo_servicio": tipo_serv,
        "direccion": dir_val,
        "estado": estado,
        "fin_contrato": fin_con,
        "num_vac": num_vac
    })

discrepancias = []
coincidencias = 0

for row in excel_rows:
    r_idx = row['row']
    proy_id = row['proyecto']
    titular_excel = row['titular']
    tit_norm = normalize_text(titular_excel)
    
    ho_proy = ho_projects.get(proy_id)
    asigs_ho = ho_asig_by_project.get(proy_id, [])
    
    disc_item = {
        "row": r_idx,
        "proyecto": proy_id,
        "servicio_excel": row['servicio'] or row['tipo_s'],
        "titular_excel": titular_excel,
        "diferencias": [],
        "propuesta": {}
    }
    
    if not ho_proy:
        disc_item["diferencias"].append(f"Proyecto '{proy_id}' no existe en SAP HO asignado a este responsable.")
    else:
        # Direccion y Cliente oficial
        den_ho = ho_proy.get('Denominacion') or ''
        cli_ho = ho_proy.get('NomCliente') or den_ho
        calle_ho = ho_proy.get('Calle') or ''
        cp_ho = ho_proy.get('CodPostal') or ''
        pob_ho = ho_proy.get('Poblacion') or ''
        dir_ho_full = f"{calle_ho}  {cp_ho}  {pob_ho}".strip()

        # Sector / Tipo de servicio oficial (ALBERGUES vs AUXILIARES vs LIMPIEZA)
        neg_ho = (ho_proy.get('NomNegocio') or '').upper()
        den_ho_u = den_ho.upper()
        cli_ho_u = cli_ho.upper()
        tipo_s_u = row['tipo_s'].upper()
        tipo_serv_excel = row['tipo_servicio'].strip().upper()

        if any(k in neg_ho or k in den_ho_u or k in cli_ho_u or k in tipo_s_u or k in tipo_serv_excel for k in ['ALBERGUE', 'XANASCAT', 'GESTBOES']):
            tipo_serv_oficial = "ALBERGUES"
        elif "LIMP" in neg_ho or "LIMP" in tipo_s_u or "LIMP" in tipo_serv_excel:
            tipo_serv_oficial = "LIMPIEZA"
        elif "AUX" in neg_ho or "CONS" in neg_ho or "AUX" in tipo_s_u or "AUX" in tipo_serv_excel:
            tipo_serv_oficial = "AUXILIARES"
        else:
            tipo_serv_oficial = row['tipo_servicio'] or "LIMPIEZA"

        if tipo_serv_excel and tipo_serv_oficial and tipo_serv_excel != tipo_serv_oficial:
            disc_item["diferencias"].append(f"Tipo Servicio: Excel '{row['tipo_servicio']}' -> Oficial '{tipo_serv_oficial}'")
            disc_item["propuesta"]["tipo_servicio_oficial"] = tipo_serv_oficial

        nombres_ho = [normalize_text(a['nombre']) for a in asigs_ho]
        
        # Titular
        if asigs_ho and tit_norm and not any(tit_norm in nho or nho in tit_norm for nho in nombres_ho):
            asig_nombres_str = ", ".join([a['nombre'] for a in asigs_ho])
            disc_item["diferencias"].append(f"Titular: Excel tiene '{titular_excel}' -> En SAP HO figura: [{asig_nombres_str}]")
            disc_item["propuesta"]["titular_sugerido"] = asigs_ho[0]['nombre']
        elif not asigs_ho and tit_norm and 'FALTA' not in tit_norm:
            disc_item["diferencias"].append(f"Titular: En SAP HO no hay operarios asignados en la bitácora activa.")

        # Dirección
        if calle_ho and normalize_text(calle_ho) not in normalize_text(row['direccion']):
            disc_item["diferencias"].append(f"Direccion: Excel '{row['direccion'][:30]}...' -> SAP HO '{dir_ho_full[:30]}...'")
            disc_item["propuesta"]["direccion_oficial"] = dir_ho_full

        # Servicio / Cliente
        if cli_ho and normalize_text(cli_ho) not in normalize_text(row['servicio']) and normalize_text(row['servicio']) not in normalize_text(cli_ho):
            disc_item["diferencias"].append(f"Cliente: Excel '{row['servicio']}' -> SAP HO '{cli_ho}'")
            disc_item["propuesta"]["cliente_oficial"] = cli_ho

        # Contrato del titular
        emp_ho = ho_employees.get(tit_norm)
        if emp_ho:
            fin_ho = emp_ho['fin_contrato']
            if row['fin_contrato'] != fin_ho and row['fin_contrato'] not in [fin_ho, '']:
                disc_item["diferencias"].append(f"Fin contrato: Excel '{row['fin_contrato']}' -> SAP HO '{fin_ho}'")
                disc_item["propuesta"]["fin_contrato_oficial"] = fin_ho
        elif tit_norm and 'FALTA' not in tit_norm:
            # Buscar en censo general si tiene otro nombre
            pass

    if disc_item["diferencias"]:
        discrepancias.append(disc_item)
    else:
        coincidencias += 1

desglose_tipologia_master = {
    "ALBERGUES": sum(1 for r in excel_rows if "ALBERGUE" in (r['tipo_servicio'] or '').upper()),
    "AUXILIARES": sum(1 for r in excel_rows if "AUX" in (r['tipo_servicio'] or '').upper()),
    "LIMPIEZA": sum(1 for r in excel_rows if "LIMP" in (r['tipo_servicio'] or '').upper())
}

print("\n" + "="*70)
print(f"RESUMEN AUDITORIA SAP HO vs MASTER EXCEL:")
print(f"Total filas analizadas: {len(excel_rows)}")
print(f"Desglose por Tipología: {desglose_tipologia_master}")
print(f"Coincidencias completas: {coincidencias}")
print(f"Filas con discrepancias detectadas: {len(discrepancias)}")
print("="*70)

for d in discrepancias:
    print(f"\n[Fila {d['row']}] {d['proyecto']} | {d['servicio_excel']}")
    print(f"   Excel: Titular='{d['titular_excel']}'")
    for diff in d['diferencias']:
        print(f"   -> {diff}")

with open('auditoria_ho_master.json', 'w', encoding='utf-8') as f:
    json.dump({
        "timestamp": datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        "total_excel": len(excel_rows),
        "desglose_tipologia": desglose_tipologia_master,
        "coincidencias": coincidencias,
        "discrepancias": discrepancias
    }, f, ensure_ascii=False, indent=2)

print("\nGuardado auditoria_ho_master.json")
